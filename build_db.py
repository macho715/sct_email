"""
Excel → DuckDB 변환 (1회 실행)
헤더: row 4, 데이터: row 5~

v2 개선 (2026-06-28) — Arrow staging table 패턴:
  - executemany UPDATE 루프 제거 (SIGSEGV 원인)
  - emb_staging에 PyArrow FixedSizeList INSERT → 단일 UPDATE JOIN
  - CHECKPOINT 주기적 WAL flush
  - 재시작 시 emb_staging 기반 자동 재개 (ON CONFLICT DO NOTHING)

v3 개선 (2026-06-29) — 청크 레벨 시맨틱 검색:
  - email_chunks 테이블: no, chunk_id, section, chunk_text, embedding FLOAT[384]
  - 이메일당 header + body_N 청크 분할 (CHUNK_SIZE=500, CHUNK_OVERLAP=50)
  - HNSW index idx_email_chunks_hnsw on email_chunks.embedding
  - app.py 청크 검색을 primary로, emails.embedding을 fallback으로 유지
"""
import re
import sys
import time
from pathlib import Path

# Windows console defaults to cp949 (Korean) which can't encode em-dash etc. in
# progress prints — force UTF-8 so the build doesn't crash before embeddings run.
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

import duckdb
import openpyxl

XLSX = Path(__file__).parent / "OUTLOOK_HVDC.xlsx"
DB   = Path(__file__).parent / "hvdc_mail.duckdb"

EMBED_MODEL   = "paraphrase-multilingual-MiniLM-L12-v2"  # 50+ langs incl. Korean, 384d
EMBED_DIM     = 384
ENCODE_BATCH  = 512    # sentence-transformers encode batch (CPU 최적)
INSERT_CHUNK  = 2048   # 한 번에 fetch·insert 할 행 수
CHUNK_SIZE    = 500    # 청크당 최대 문자 수
CHUNK_OVERLAP = 50     # 청크 간 겹침 문자 수
MAX_BODY_CHUNKS = 10   # 이메일당 최대 body 청크 수


def clean_col(name: str) -> str:
    if not name:
        return "col_unknown"
    return re.sub(r"\W+", "_", str(name)).strip("_").lower()


# ══════════════════════════════════════════════════════════ Phase 1 ══
def build() -> None:
    """Excel → emails 테이블 + FTS + B-tree 인덱스."""
    print(f"Loading {XLSX.name} ...")
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb["전체_데이터"]
    rows_iter = ws.iter_rows(values_only=True)

    for _ in range(3):            # rows 1-3 skip (title/note/blank)
        next(rows_iter)
    raw_headers = next(rows_iter) # row 4 = headers
    headers = [clean_col(h) for h in raw_headers]

    print(f"Columns ({len(headers)}): {headers}")
    data = [list(r) for r in rows_iter if any(v is not None for v in r)]
    print(f"  {len(data):,} data rows loaded")

    con = duckdb.connect(str(DB))
    con.execute("DROP TABLE IF EXISTS emails")
    col_defs = ", ".join(f'"{h}" VARCHAR' for h in headers)
    con.execute(f"CREATE TABLE emails ({col_defs})")

    CHUNK = 5_000
    ph = ", ".join(["?"] * len(headers))
    for i in range(0, len(data), CHUNK):
        con.executemany(
            f'INSERT INTO emails VALUES ({ph})',
            [list(map(str, r)) for r in data[i : i + CHUNK]],
        )
        print(f"  inserted {min(i + CHUNK, len(data)):,}/{len(data):,}", end="\r")

    print("\nBuilding FTS index ...")
    con.execute("INSTALL fts; LOAD fts;")
    con.execute("""
        PRAGMA create_fts_index(
            'emails', 'no',
            'subject', 'sendername', 'senderemail',
            'recipientto', 'plaintextbody',
            'hvdc_cases', 'company_name',
            stemmer='english', stopwords='english',
            ignore='(\\.|[^a-z])+',
            strip_accents=1, lower=1, overwrite=1
        )
    """)
    con.execute('CREATE INDEX IF NOT EXISTS idx_month ON emails("month")')
    con.execute('CREATE INDEX IF NOT EXISTS idx_stage ON emails("stage")')
    con.execute('CREATE INDEX IF NOT EXISTS idx_site  ON emails("site")')

    n = con.execute("SELECT COUNT(*) FROM emails").fetchone()[0]
    print(f"Phase 1 done — {n:,} rows in {DB.name}")
    con.close()


# ══════════════════════════════════════════════════════════ helpers ══
def _drop_indexes(con: duckdb.DuckDBPyConnection) -> None:
    for sql in [
        "DROP INDEX IF EXISTS idx_month",
        "DROP INDEX IF EXISTS idx_stage",
        "DROP INDEX IF EXISTS idx_site",
        "DROP INDEX IF EXISTS idx_embedding_hnsw",
    ]:
        try:
            con.execute(sql)
        except Exception:
            pass
    # NOTE: do NOT drop_fts_index here — it deletes the 'stopwords' dependency and
    # poisons the session, making the later create_fts_index fail on commit
    # ("subject stopwords has been deleted"). create_fts_index(overwrite=1) in
    # _rebuild_indexes already replaces any existing index cleanly.


def _rebuild_indexes(con: duckdb.DuckDBPyConnection) -> None:
    con.execute('CREATE INDEX IF NOT EXISTS idx_month ON emails("month")')
    con.execute('CREATE INDEX IF NOT EXISTS idx_stage ON emails("stage")')
    con.execute('CREATE INDEX IF NOT EXISTS idx_site  ON emails("site")')
    print("  B-tree indexes OK.")

    try:
        con.execute("INSTALL vss; LOAD vss;")
        con.execute("SET hnsw_enable_experimental_persistence = true")
        con.execute("DROP INDEX IF EXISTS idx_embedding_hnsw")
        con.execute(f"""
            CREATE INDEX idx_embedding_hnsw
            ON emails USING HNSW (embedding)
            WITH (metric = 'cosine')
        """)
        print("  HNSW index created.")
    except Exception as exc:
        print(f"  HNSW skipped: {exc}")

    try:
        con.execute("INSTALL fts; LOAD fts;")
        con.execute("""
            PRAGMA create_fts_index(
                'emails', 'no',
                'subject', 'sendername', 'senderemail',
                'recipientto', 'plaintextbody',
                'hvdc_cases', 'company_name',
                stemmer='english', stopwords='english',
                ignore='(\\.|[^a-z])+',
                strip_accents=1, lower=1, overwrite=1
            )
        """)
        print("  FTS index rebuilt.")
    except Exception as exc:
        print(f"  FTS skipped: {exc}")


# ══════════════════════════════════════════════════════════ Phase 2 ══
def add_embeddings(db_path: Path) -> None:
    """
    Arrow staging table 패턴으로 FLOAT[384] 임베딩 추가.

    흐름:
      1. emb_staging(no PK, embedding) 테이블에 PyArrow INSERT 누적
         → ON CONFLICT DO NOTHING으로 재시작 시 중복 스킵
      2. 모든 배치 완료 후 단일 UPDATE JOIN emails ← emb_staging
      3. emb_staging DROP + CHECKPOINT
      4. HNSW + FTS + B-tree 인덱스 재생성

    executemany UPDATE 루프를 사용하지 않으므로 WAL 과부하·SIGSEGV 없음.
    """
    import numpy as np
    import pyarrow as pa
    from sentence_transformers import SentenceTransformer

    print(f"Loading {EMBED_MODEL} ...")
    model = SentenceTransformer(EMBED_MODEL)

    con = duckdb.connect(str(db_path))
    total = con.execute("SELECT COUNT(*) FROM emails").fetchone()[0]

    # ── 재시작 감지 ─────────────────────────────────────────────────
    staging_exists = bool(
        con.execute(
            "SELECT 1 FROM information_schema.tables WHERE table_name = 'emb_staging'"
        ).fetchone()
    )
    staging_count = (
        con.execute("SELECT COUNT(*) FROM emb_staging").fetchone()[0]
        if staging_exists else 0
    )

    if staging_count > 0:
        print(f"  Resume: emb_staging already has {staging_count:,} rows. Continuing ...")
    else:
        print("  Fresh run: dropping indexes and resetting embedding column.")
        _drop_indexes(con)
        con.execute("ALTER TABLE emails DROP COLUMN IF EXISTS embedding")
        con.execute(f"ALTER TABLE emails ADD COLUMN embedding FLOAT[{EMBED_DIM}]")
        con.execute("DROP TABLE IF EXISTS emb_staging")
        con.execute(f"""
            CREATE TABLE emb_staging (
                "no"       VARCHAR PRIMARY KEY,
                embedding  FLOAT[{EMBED_DIM}]
            )
        """)

    # ── Arrow batch INSERT 루프 ─────────────────────────────────────
    #
    # DuckDB 공식 권장: "avoid using lots of individual row-by-row INSERT
    # statements" — Arrow 등록을 사용해 zero-copy bulk INSERT.
    # ref: https://duckdb.org/docs/guides/python/import_arrow
    #
    fixed_list_type = pa.list_(pa.float32(), EMBED_DIM)  # FixedSizeList<384 x float32>
    processed = staging_count
    t0 = time.time()

    print(f"Encoding {total:,} rows "
          f"(ENCODE_BATCH={ENCODE_BATCH}, INSERT_CHUNK={INSERT_CHUNK}) ...")

    while True:
        rows = con.execute(
            """
            SELECT e."no", e.subject, e.plaintextbody
            FROM   emails e
            WHERE  NOT EXISTS (
                       SELECT 1 FROM emb_staging s WHERE s."no" = e."no"
                   )
            LIMIT  ?
            """,
            [INSERT_CHUNK],
        ).fetchall()

        if not rows:
            break

        row_nos = [r[0] for r in rows]
        texts   = [f"{r[1] or ''} {(r[2] or '')[:500]}".strip() for r in rows]

        # sentence-transformers encode
        # convert_to_numpy=True → numpy float32 배열 직접 반환 (list 변환 없음)
        emb_np = model.encode(
            texts,
            batch_size=ENCODE_BATCH,
            normalize_embeddings=True,
            show_progress_bar=False,
            convert_to_numpy=True,
        ).astype("float32")   # shape (N, EMBED_DIM)

        # numpy (N, 384) → PyArrow FixedSizeListArray
        flat = emb_np.flatten()   # (N * 384,) contiguous float32
        emb_arrow = pa.FixedSizeListArray.from_arrays(
            pa.array(flat, type=pa.float32()),
            list_size=EMBED_DIM,
        )
        arrow_tbl = pa.table({
            "no":        pa.array(row_nos, type=pa.string()),
            "embedding": emb_arrow,
        })

        # DuckDB에 Arrow 테이블 등록 → 명시적 컬럼 매핑으로 컬럼 스왑 버그 방지
        con.register("_emb_view", arrow_tbl)
        con.execute(
            'INSERT INTO emb_staging ("no", embedding) '
            'SELECT "no", embedding FROM _emb_view '
            'ON CONFLICT DO NOTHING'
        )
        con.unregister("_emb_view")

        processed += len(rows)

        if processed % 5_000 < INSERT_CHUNK:
            elapsed = time.time() - t0
            rate    = processed / elapsed if elapsed > 0 else 1
            eta_min = (total - processed) / rate / 60
            print(
                f"  {processed:,}/{total:,} | {rate:.0f} rows/s | ETA {eta_min:.1f}min",
                flush=True,
            )

        # WAL 주기적 flush → 파일 크기 안정화
        if processed % 10_000 < INSERT_CHUNK:
            con.execute("CHECKPOINT")
            print(f"  [CHECKPOINT] {processed:,} rows flushed to disk", flush=True)

    # ── 단일 UPDATE JOIN (emb_staging → emails) ─────────────────────
    staged = con.execute("SELECT COUNT(*) FROM emb_staging").fetchone()[0]
    print(f"\nFinal UPDATE: {staged:,} staged embeddings → emails.embedding ...")
    con.execute("""
        UPDATE emails
        SET    embedding = s.embedding
        FROM   emb_staging s
        WHERE  emails."no" = s."no"
    """)
    con.execute("DROP TABLE emb_staging")
    con.execute("CHECKPOINT")

    done = con.execute(
        "SELECT COUNT(*) FROM emails WHERE embedding IS NOT NULL"
    ).fetchone()[0]
    print(f"Embedded: {done:,} / {total:,}")

    # ── 인덱스 재생성 ────────────────────────────────────────────────
    _rebuild_indexes(con)
    con.execute("CHECKPOINT")
    con.close()
    print("Phase 2 done.")


# ══════════════════════════════════════════════════════════ Phase 3 ══

def _split_text_into_chunks(text: str) -> list[str]:
    """Split text into overlapping CHUNK_SIZE-char chunks."""
    text = (text or "").strip()
    if not text:
        return []
    if len(text) <= CHUNK_SIZE:
        return [text]
    chunks = []
    start = 0
    step = max(1, CHUNK_SIZE - CHUNK_OVERLAP)
    while start < len(text):
        chunks.append(text[start : start + CHUNK_SIZE])
        start += step
    return chunks


def build_chunks(db_path: Path) -> None:
    """
    emails 테이블 → email_chunks 테이블 생성 + HNSW 인덱스.

    각 이메일을 두 가지 section으로 분할:
      header  : subject | sendername | company_name (메타데이터 청크)
      body_N  : plaintextbody를 CHUNK_SIZE 단위로 분할 (최대 MAX_BODY_CHUNKS개)

    chunk_id = "{no}_header" | "{no}_body_{N}"
    ON CONFLICT DO NOTHING 으로 재시작 시 중복 스킵.

    완료 후:
      - HNSW index idx_email_chunks_hnsw 생성
      - EXPLAIN SELECT ... ORDER BY cosine_score DESC LIMIT 10 으로 HNSW_INDEX_SCAN 확인
    """
    import numpy as np
    import pyarrow as pa
    from sentence_transformers import SentenceTransformer

    print(f"Loading {EMBED_MODEL} for chunk embeddings ...")
    model = SentenceTransformer(EMBED_MODEL)

    con = duckdb.connect(str(db_path))

    # ── 테이블 생성 (이미 있으면 재시작으로 간주) ───────────────────
    table_exists = bool(
        con.execute(
            "SELECT 1 FROM information_schema.tables WHERE table_name = 'email_chunks'"
        ).fetchone()
    )
    if not table_exists:
        con.execute(f"""
            CREATE TABLE email_chunks (
                "no"       VARCHAR,
                chunk_id   VARCHAR PRIMARY KEY,
                section    VARCHAR,
                chunk_text VARCHAR,
                embedding  FLOAT[{EMBED_DIM}]
            )
        """)
        print("  email_chunks table created.")
    else:
        already = con.execute("SELECT COUNT(*) FROM email_chunks").fetchone()[0]
        print(f"  email_chunks already has {already:,} chunks — resuming.")

    total_emails = con.execute("SELECT COUNT(*) FROM emails").fetchone()[0]
    print(f"Chunking {total_emails:,} emails ...")

    fixed_list_type = pa.list_(pa.float32(), EMBED_DIM)
    processed = 0
    t0 = time.time()

    while True:
        rows = con.execute(
            f"""
            SELECT "no", subject, sendername, company_name, plaintextbody
            FROM   emails
            WHERE  "no" NOT IN (
                       SELECT DISTINCT "no" FROM email_chunks
                   )
            LIMIT  {INSERT_CHUNK}
            """
        ).fetchall()

        if not rows:
            break

        chunk_nos      = []
        chunk_ids      = []
        chunk_sections = []
        chunk_texts    = []

        for row in rows:
            no, subject, sendername, company_name, body = row

            # header chunk
            header_parts = [
                str(subject      or "").strip(),
                str(sendername   or "").strip(),
                str(company_name or "").strip(),
            ]
            header_text = " | ".join(p for p in header_parts if p)
            if header_text:
                chunk_nos.append(no)
                chunk_ids.append(f"{no}_header")
                chunk_sections.append("header")
                chunk_texts.append(header_text[:CHUNK_SIZE])

            # body chunks
            body_parts = _split_text_into_chunks(str(body or ""))
            for i, part in enumerate(body_parts[:MAX_BODY_CHUNKS]):
                chunk_nos.append(no)
                chunk_ids.append(f"{no}_body_{i}")
                chunk_sections.append(f"body_{i}")
                chunk_texts.append(part)

        if not chunk_texts:
            break

        # encode
        emb_np = model.encode(
            chunk_texts,
            batch_size=ENCODE_BATCH,
            normalize_embeddings=True,
            show_progress_bar=False,
            convert_to_numpy=True,
        ).astype("float32")

        flat = emb_np.flatten()
        emb_arrow = pa.FixedSizeListArray.from_arrays(
            pa.array(flat, type=pa.float32()),
            list_size=EMBED_DIM,
        )
        arrow_tbl = pa.table({
            "no":         pa.array(chunk_nos,      type=pa.string()),
            "chunk_id":   pa.array(chunk_ids,      type=pa.string()),
            "section":    pa.array(chunk_sections, type=pa.string()),
            "chunk_text": pa.array(chunk_texts,    type=pa.string()),
            "embedding":  emb_arrow,
        })

        con.register("_chunk_view", arrow_tbl)
        con.execute(
            'INSERT INTO email_chunks ("no", chunk_id, section, chunk_text, embedding) '
            'SELECT "no", chunk_id, section, chunk_text, embedding FROM _chunk_view '
            'ON CONFLICT (chunk_id) DO NOTHING'
        )
        con.unregister("_chunk_view")

        processed += len(rows)
        if processed % 5_000 < INSERT_CHUNK:
            elapsed = time.time() - t0
            rate    = processed / elapsed if elapsed > 0 else 1
            eta_min = (total_emails - processed) / rate / 60
            print(
                f"  {processed:,}/{total_emails:,} | {rate:.0f} emails/s | ETA {eta_min:.1f}min",
                flush=True,
            )

        if processed % 10_000 < INSERT_CHUNK:
            con.execute("CHECKPOINT")
            print(f"  [CHECKPOINT] {processed:,} emails flushed.", flush=True)

    total_chunks = con.execute("SELECT COUNT(*) FROM email_chunks").fetchone()[0]
    print(f"\nTotal chunks: {total_chunks:,}")
    con.execute("CHECKPOINT")

    # ── HNSW 인덱스 생성 ────────────────────────────────────────────
    try:
        con.execute("INSTALL vss; LOAD vss;")
        con.execute("SET hnsw_enable_experimental_persistence = true")
        con.execute("DROP INDEX IF EXISTS idx_email_chunks_hnsw")
        con.execute(f"""
            CREATE INDEX idx_email_chunks_hnsw
            ON email_chunks USING HNSW (embedding)
            WITH (metric = 'cosine')
        """)
        print("  HNSW index idx_email_chunks_hnsw created.")

        # EXPLAIN으로 HNSW_INDEX_SCAN 확인
        dummy_vec = [0.0] * EMBED_DIM
        explain_rows = con.execute(
            f"EXPLAIN SELECT \"no\", chunk_id, "
            f"array_cosine_similarity(embedding, ?::FLOAT[{EMBED_DIM}]) AS cosine_score "
            f"FROM email_chunks "
            f"ORDER BY cosine_score DESC LIMIT 10",
            [dummy_vec],
        ).fetchall()
        explain_text = "\n".join(" ".join(str(c) for c in r) for r in explain_rows)
        if "HNSW_INDEX_SCAN" in explain_text:
            print("  ✓ HNSW_INDEX_SCAN confirmed in EXPLAIN output.")
        else:
            print("  ⚠ HNSW_INDEX_SCAN not found — index may not be used for this query.")
    except Exception as exc:
        print(f"  HNSW index skipped: {exc}")

    con.execute("CHECKPOINT")
    con.close()
    print("Phase 3 (chunks) done.")


if __name__ == "__main__":
    build()
    add_embeddings(DB)
    build_chunks(DB)
